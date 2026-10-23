import copy
import io
import json
import re
import zipfile
from datetime import date, datetime
from pathlib import Path
from xml.sax.saxutils import escape

import pandas as pd
import streamlit as st
from PIL import Image, ImageOps

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (Image as RLImage, KeepTogether, Paragraph,
                                SimpleDocTemplate, Spacer, Table, TableStyle)

st.set_page_config(page_title="Generator Dokumen Translok BPS", page_icon="📝", layout="wide")

ASSETS = Path(__file__).parent / "assets"
KOP_PATH = ASSETS / "kop.png"
QR_PATH = ASSETS / "qr.png"
KOP_WIDTH_CM = 10.5

# ----------------------------------------------------------------------------
# DATA
# ----------------------------------------------------------------------------
SURVEI = [
    "Survei Kerangka Sampel Area (Survei KSA)", "Survei Ubinan", "SP Palawija", "SP Lahan",
    "Survei Konversi Gabah ke Beras (SKGB)", "Statistik Pertanian Hortikultura (SPH)", "VN Horti",
    "Survei Indeks Produksi Hortikultura (VIP-Hortikultura)",
    "Survei Perkebunan Bulanan (SEDAPP Online)", "Survei Perusahaan Perkebunan (SKB) Tahunan",
    "Updating Direktori Perusahaan Pertanian dan Usaha Tani Lainnya (UDPP DUTL)",
    "Survei Komoditas Strategis (Komstrat)", "Laporan Pemotongan Ternak Bulanan (LPTB)",
    "Laporan Tahunan Perusahaan Budidaya dan Penangkapan Ikan (LTB dan LTP)", "Survei Kehutanan",
    "Pencacahan Lengkap Tempat Pelelangan Ikan Triwulanan (PP/TPI)",
    "Survei IBS Tahunan - Pemutakhiran Direktori Perusahaan Awal (DPA)",
    "Survei IBS Tahunan - Survei Tahunan Perusahaan Industri Manufaktur - STPIM A",
    "Survei IBS Triwulan B", "STPU (Survei Triwulanan Pelaku Usaha) Triwulan B",
    "Survei IMK Triwulan B", "Survei IMK Tahunan", "Updating Perusahaan PE (Pertambangan Energi)",
    "Survei PE Tahunan (Survei Galian URT)", "Survei PE Tahunan (Survei Non Migas)",
    "Survei PE Tahunan (Survei Captive Power)", "Survei PE Tahunan (Survei Air Bersih)",
    "Survei PE Triwulanan (Air Bersih) Triwulan B", "Pengutipan Direktori Konstruksi",
    "Survei Konstruksi Tahunan", "Survei Kesejahteraan Petani (SKP)",
    "Survei Konstruksi Triwulanan Triwulan B",
]
KEGIATAN = ["Listing", "Pendataan"]
TAHUN = [2026, 2027, 2028, 2029, 2030]
SUBROUND = ["Subround 1 (SR 1)", "Subround 2 (SR 2)", "Subround 3 (SR 3)"]
JENIS_UBINAN = ["Palawija", "Padi"]

KECAMATAN = {
    "Payung": "PAYUNG", "Pulau Besar": "PULAU BESAR", "Simpang Rimba": "SIMPANG RIMBA",
    "Toboali": "TOBOALI", "Tukak Sadai": "TUKAK SADAI", "Airgegas": "AIR GEGAS",
    "Lepar": "LEPAR PONGOK", "Kepulauan Pongok": "KEPULAUAN PONGOK",
}
KEC_DOC = {"Airgegas": "Air Gegas"}

_DESA_RAW = """Desa GADUNG|TOBOALI
Kelurahan TELADAN|TOBOALI
Kelurahan TANJUNG KETAPANG|TOBOALI
Kelurahan TOBOALI|TOBOALI
Desa SERDANG|TOBOALI
Desa JERIJI|TOBOALI
Desa KEPOH|TOBOALI
Desa BIKANG|TOBOALI
Desa RINDIK|TOBOALI
Desa RIAS|TOBOALI
Desa KAPOSANG|TOBOALI
Desa SEBAGIN|SIMPANG RIMBA
Desa GUDANG|SIMPANG RIMBA
Desa SIMPANG RIMBA|SIMPANG RIMBA
Desa BANGKA KOTA|SIMPANG RIMBA
Desa RAJIK|SIMPANG RIMBA
Desa PERMIS|SIMPANG RIMBA
Desa JELUTUNG II|SIMPANG RIMBA
Desa MALIK|PAYUNG
Desa NADUNG|PAYUNG
Desa RANGGUNG|PAYUNG
Desa IRAT|PAYUNG
Desa PAYUNG|PAYUNG
Desa BEDENGUNG|PAYUNG
Desa PAKU|PAYUNG
Desa PANGKALBULUH|PAYUNG
Desa SENGIR|PAYUNG
Desa FAJAR INDAH|PULAU BESAR
Desa SUMBER JAYA PERMAI|PULAU BESAR
Desa PANCA TUNGGAL|PULAU BESAR
Desa SUKAJAYA|PULAU BESAR
Desa BATU BETUMPANG|PULAU BESAR
Desa NANGKA|AIR GEGAS
Desa TEPUS|AIR GEGAS
Desa DELAS|AIR GEGAS
Desa NYELANDING|AIR GEGAS
Desa SIDOHARJO|AIR GEGAS
Desa PERGAM|AIR GEGAS
Desa AIRGEGAS|AIR GEGAS
Desa BENCAH|AIR GEGAS
Desa AIR BARA|AIR GEGAS
Desa RANGGAS|AIR GEGAS
Desa TIRAM|TUKAK SADAI
Desa BUKIT TERAP|TUKAK SADAI
Desa TUKAK|TUKAK SADAI
Desa SADAI|TUKAK SADAI
Desa PASIR PUTIH|TUKAK SADAI
Desa CELAGEN|KEPULAUAN PONGOK
Desa PONGOK|KEPULAUAN PONGOK
Desa KUMBUNG|LEPAR PONGOK
Desa PENUTUK|LEPAR PONGOK
Desa TANJUNG SANGKAR|LEPAR PONGOK
Desa TANJUNG LABU|LEPAR PONGOK"""
DESA_DB = dict(line.split("|") for line in _DESA_RAW.splitlines())

# nama | NIP | pangkat | gol | jabatan
_PEG_RAW = """Agung Rachmadi, SE|197111261999031003|Pembina Tk I|IVb|Kepala BPS Kabupaten Bangka Selatan
Imam Hidayat, SST, M.Si|197807081999121001|Pembina|IVa|Kepala Subbagian Umum
Damri Ismail|197102212002121003|Penata Tk I|IIId|Statistisi Penyelia
Firman, SST., M.M.|198604052009121002|Pembina|IVa|Statistisi Ahli Madya
KMS M Hidayat, SST|198902062013111001|Penata Tk I|IIId|Pranata Komputer Ahli Muda
Sridewi Nainggolan, S.Si, M.Si|198505132009022009|Penata Tk I|IIId|Statistisi Ahli Muda
Rezky Mayda Putri, SST|199305212016022001|Penata|IIIc|Statistisi Ahli Muda
Ali Shaumanjaya, A.Md|197409182002121004|Penata Muda Tk I|IIIb|Statistisi Mahir
Zullizar, SE|198510092010031003|Penata Tk I|IIId|Analis Pengelolaan Keuangan APBN Ahli Muda
Raka Suryadi, S.Stat|199907012019121001|Penata Muda|IIIa|Pranata Keuangan APBN Terampil
Sutri Vininda, S.Tr. Stat|199903122021042001|Penata Muda Tk I|IIIb|Statistisi Ahli Pertama
Risma Dillah, A.Md.Stat|200104132022012001|Pengatur Tk I|IId|Statistisi Terampil
Adinda Hermambang, S.Tr.Stat|199812302022012002|Penata Muda Tk I|IIIb|Statistisi Ahli Pertama
Zahrotul Firdaus, S.Tr.Stat|199902182022012001|Penata Muda Tk I|IIIb|Statistisi Ahli Pertama
Ragil Chelvin Pratama, A.Md.Stat|200112212023021002|Pengatur|IIc|Statistisi Terampil
Fasyina Tsurayya Nisyah, S.Tr.Stat|200009122023102001|Penata Muda|IIIa|Statistisi Ahli Pertama
Prasetya Dwi Anggara Kushendrawan, SST|199410172018021001|Penata Muda Tk I|IIIb|Statistisi Ahli Pertama
Ruri Tri Wahyuni, S.Tr.Stat|200210092024122001|Penata Muda|IIIa|Statistisi Ahli Pertama
Ego, A.Md.Stat|200202072024121003|Pengatur|IIc|Statistisi Terampil
Nurusyifa Somanaya, S.Tr.Stat|200309202024122001|Penata Muda|IIIa|Statistisi Ahli Pertama
Maya Imas Tirany, S.Tr.Stat|200007172024122001|Penata Muda|IIIa|Statistisi Ahli Pertama
Lailatul Hasanah, S.Tr.Stat|200112032024122001|Penata Muda|IIIa|Pranata Komputer Ahli Pertama
Umi Syafiyah, S.Tr.Stat|200107222023102002|Penata Muda|IIIa|Statistisi Ahli Pertama
Dwi Lestari, S.Tr. Stat|199811072019122001|Penata Muda|IIIa|Statistisi Ahli Pertama
Stenislaus Angga Aprianto, S.Tr.Stat.|200304072026031003|Penata Muda|IIIa|Staf
Satrya, A.Md|198601272024211001||VII|Pranata Komputer Terampil
Rosidi|198703232025211051||V|Operator Layanan Operasional
Samsudin|198704202025211050||V|Operator Layanan Operasional"""


def _clean_name(n):
    n = n.replace("S.Tr. Stat", "S.Tr.Stat")
    return n[:-1] if n.endswith("S.Tr.Stat.") else n


PEGAWAI = []
for _l in _PEG_RAW.splitlines():
    a, b, c, d_, e = _l.split("|")
    PEGAWAI.append(dict(nama=_clean_name(a), nip=b, pangkat=c, gol=d_, jabatan=e))
PELAKSANA = [(p["nama"], p["nip"]) for p in PEGAWAI]

BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
         "Agustus", "September", "Oktober", "November", "Desember"]
ROMAN = {"I", "II", "III", "IV", "V", "VI"}

PAPER = {"A4": (21.0, 29.7), "F4 / Folio": (21.5, 33.0), "Letter": (21.59, 27.94)}
MARGIN_LR = 2.5

POS_OPTIONS = ["Tetap pada urutan semestinya (sebelum Dokumentasi)",
               "Sisipkan setelah poin 1", "Sisipkan setelah poin 2", "Sisipkan setelah poin 3"]

DEFAULT_STYLE = {
    "Listing": {
        "perihal": "Laporan Perjalanan Dinas Dalam Kota dalam rangka melakukan Pengawasan {kegiatan} {survei_tahun}",
        "p1": "Pada saat melakukan pengawasan {kegiatan} {survei_tahun} di Kecamatan {kecamatan} tanggal {hari} {bulan} tahun {tahun_pelaksanaan}, saya bertemu dengan PCL {pcl}.",
        "p2": "Saya melakukan pengawasan {kegiatan} {survei_tahun} di 1 SLS yaitu {sls}, {desa} dan bertemu dengan beberapa responden, yaitu {responden}, dll.",
        "p3": "Pengawasan dilakukan dengan mendampingi dan mengawasi petugas lapangan saat melakukan listing {survei} door to door di SLS tersebut, dan melihat ketepatan wawancaranya.",
    },
    "Pendataan": {
        "perihal": "Laporan Perjalanan Dinas Dalam Kota dalam rangka melakukan Pengawasan {kegiatan} {survei_tahun}",
        "p1": "Pada saat melakukan pengawasan {kegiatan} {survei_tahun} di Kecamatan {kecamatan} tanggal {hari} {bulan} tahun {tahun_pelaksanaan}, saya bertemu dengan PCL {pcl}.",
        "p2": "Saya melakukan pengawasan {kegiatan} {survei_tahun} di 1 sampel, yaitu {sls} {desa} dan bertemu dengan Penanggungjawab Usaha, yaitu {responden}.",
        "p3": "Pengawasan dilakukan dengan mendampingi dan mengawasi petugas lapangan saat melakukan {kegiatan} {survei} di desa tersebut, dan melihat ketepatan wawancara dan isiannya.",
    },
}
BASE_VARS = ["kegiatan", "survei", "survei_tahun", "tahun", "kecamatan", "desa", "sls", "pcl",
             "responden", "tanggal", "hari", "bulan", "tahun_pelaksanaan", "pelaksana"]

DEFAULT_COLS = ["Kegiatan", "Survei", "Tahun", "Subround", "Jenis Ubinan", "Triwulan", "Kecamatan",
                "Desa", "Nama SLS/Sampel", "Nama PCL", "Nama Responden", "Tanggal Pelaksanaan",
                "Nama Pelaksana", "NIP Pelaksana", "Nama File Foto"]

SUPER_P2 = ("Demikian pernyataan ini kami buat dengan sebenar-benarnya untuk dipergunakan sebagaimana "
            "mestinya. Apabila terdapat kekeliruan dalam pertanggungjawaban SPD dan mengakibatkan "
            "kerugian negara, saya bersedia dituntut sesuai aturan yang berlaku dan mengembalikan "
            "biaya transport yang sudah saya terima ke kas negara.")


# ----------------------------------------------------------------------------
# HELPERS
# ----------------------------------------------------------------------------
def pretty_desa(raw):
    return " ".join(w if w in ROMAN else w.capitalize() for w in raw.split())


def tgl_id(d):
    return f"{d.day} {BULAN[d.month - 1]} {d.year}"


def survei_label(survei, tahun, triwulan, subround, jenis):
    if survei == "Survei Ubinan":
        return f"Survei Ubinan {jenis} {subround}", f" tahun {tahun}"
    label, suffix = survei, f" tahun {tahun}"
    if re.search(r"\bA$", label):
        label, suffix = re.sub(r"\bA$", str(tahun), label), ""
    if re.search(r"\bB$", label):
        label = re.sub(r"\bB$", str(triwulan), label)
    return label, suffix


def has_B(survei):
    return bool(re.search(r"\bB$", survei))


def split_names(text):
    return [x.strip() for x in re.split(r"[,\n;]+", text or "") if x.strip()]


def render_tpl(tpl, v):
    return re.sub(r"\{(\w+)\}", lambda m: str(v.get(m.group(1), m.group(0))), tpl)


def var_key(name):
    return re.sub(r"\W+", "_", str(name).strip().lower()).strip("_")


def sign_name(n):
    return n + "." if n.endswith("Stat") else n


@st.cache_data(show_spinner=False)
def process_photo_bytes(data: bytes):
    img = ImageOps.exif_transpose(Image.open(io.BytesIO(data))).convert("RGB")
    img.thumbnail((1400, 1400))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=85)
    w, h = img.size
    w_cm = 8.5
    h_cm = w_cm * h / w
    if h_cm > 9.5:
        h_cm = 9.5
        w_cm = h_cm * w / h
    return buf.getvalue(), w_cm, h_cm


@st.cache_data(show_spinner=False)
def to_gray_bytes(data: bytes):
    buf = io.BytesIO()
    Image.open(io.BytesIO(data)).convert("L").save(buf, "JPEG", quality=85)
    return buf.getvalue()


def scale_photos(photos, width_cm, max_h_cm, paper="A4"):
    """Atur ukuran foto: lebar = width_cm, tinggi dibatasi max_h_cm, tidak melebihi lebar halaman."""
    cw = PAPER[paper][0] - 2 * MARGIN_LR
    out = []
    for b, w, h in photos:
        ratio = h / w
        nw = min(width_cm, cw)
        nh = nw * ratio
        if nh > max_h_cm:
            nh = max_h_cm
            nw = nh / ratio
        if st.session_state.get("photo_color") == "Grayscale (hitam-putih)":
            b = to_gray_bytes(b)
        out.append((b, nw, nh))
    return out


def make_vars(rec):
    label, suffix = survei_label(rec["survei"], rec["tahun"], rec.get("triwulan"),
                                 rec.get("subround"), rec.get("jenis"))
    t = rec["tanggal"]
    v = dict(kegiatan=rec["kegiatan"], survei=label, survei_tahun=label + suffix,
             tahun=str(rec["tahun"]), kecamatan=KEC_DOC.get(rec["kec"], rec["kec"]),
             desa=display_desa(rec["desa"]), sls=rec["sls"], pcl=rec["pcl"],
             responden=", ".join(rec["responden"]), tanggal=tgl_id(t), hari=str(t.day),
             bulan=BULAN[t.month - 1], tahun_pelaksanaan=str(t.year),
             pelaksana=rec["pelaksana_nama"])
    v.update(rec.get("extra", {}))
    return v


def assemble_points(texts, extras, v):
    out = []
    for k in (1, 2, 3):
        out.append(texts[k - 1])
        out += [render_tpl(e["text"], v) for e in extras if e["after"] == k]
        if k == 3:
            out += [render_tpl(e["text"], v) for e in extras if e["after"] == 0]
    return out


def doc_data(rec, v, texts_perihal, points, photos, kop):
    nama, nip = rec["pelaksana_nama"], rec["pelaksana_nip"]
    return {"kop": kop, "kepada": "Kepala BPS Kabupaten Bangka Selatan", "perihal": texts_perihal,
            "tujuan": f"Kecamatan {v['kecamatan']}", "tanggal": v["tanggal"], "points": points,
            "photos": photos, "nama": nama, "nip": nip}


# ----------------------------------------------------------------------------
# WORD
# ----------------------------------------------------------------------------
def _font(run, size=12, bold=False, italic=False, underline=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold, run.italic, run.underline = bold, italic, underline


ITALIC_RE = re.compile(r"(Subround|door to door)")


def _runs(par, text, size=12):
    for part in ITALIC_RE.split(text):
        if part:
            _font(par.add_run(part), size, italic=bool(ITALIC_RE.fullmatch(part)))


def _fmt(par, after=0, justify=False):
    pf = par.paragraph_format
    pf.space_before, pf.space_after = Pt(0), Pt(after)
    pf.left_indent, pf.first_line_indent = Cm(0), Cm(0)  # tanpa menjorok / menggantung
    pf.widow_control = True
    if justify:
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _setup_doc(paper):
    doc = Document()
    sec = doc.sections[0]
    w, h = PAPER[paper]
    sec.page_width, sec.page_height = Cm(w), Cm(h)
    sec.left_margin = sec.right_margin = Cm(MARGIN_LR)
    sec.top_margin = sec.bottom_margin = Cm(2)
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)
    return doc, sec, w - 2 * MARGIN_LR


def _kv_table(doc, rows, widths):
    t = doc.add_table(rows=len(rows), cols=3)
    t.autofit = False
    for i, (a, b) in enumerate(rows):
        for j, txt in enumerate([a, ":", b]):
            c = t.rows[i].cells[j]
            c.width = Cm(widths[j])
            _fmt(c.paragraphs[0])
            _runs(c.paragraphs[0], txt)


def build_docx(d, paper="A4"):
    doc, sec, cw = _setup_doc(paper)
    if d["kop"] and KOP_PATH.exists():
        p = doc.add_paragraph()
        _fmt(p, 14)
        p.add_run().add_picture(str(KOP_PATH), width=Cm(KOP_WIDTH_CM))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(p.add_run("Laporan Perjalanan Translok"), 13, bold=True, underline=True)
    _fmt(p, 20)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _kv_table(doc, [("Kepada", d["kepada"]), ("Perihal", d["perihal"]), ("Tujuan", d["tujuan"]),
                    ("Tanggal Pelaksanaan", d["tanggal"])], [4.2, 0.4, cw - 4.6])
    sp = doc.add_paragraph()
    _fmt(sp, 6)
    p = doc.add_paragraph()
    _runs(p, "Hasil Pelaksanaan")
    _fmt(p, 2)
    n = 0
    for txt in [p for p in d["points"] if p.strip()] + ["Dokumentasi:"]:
        n += 1
        par = doc.add_paragraph()
        _fmt(par, 3, justify=True)
        par.paragraph_format.keep_together = True
        _runs(par, f"{n}. {txt}")
    for img_bytes, w_cm, h_cm in d["photos"]:
        par = doc.add_paragraph()
        _fmt(par, 4)
        par.add_run().add_picture(io.BytesIO(img_bytes), width=Cm(w_cm), height=Cm(h_cm))
    sp = doc.add_paragraph()
    _fmt(sp, 6)
    sig = doc.add_table(rows=1, cols=2)
    sig.autofit = False
    sig.alignment = WD_TABLE_ALIGNMENT.LEFT
    sig.rows[0].cells[0].width = Cm(cw - 7.5)
    sig.rows[0].cells[1].width = Cm(7.5)
    cell = sig.rows[0].cells[1]
    lines = [f"Toboali, {d['tanggal']}", "Pelaksana SPD", "", "", "", "", d["nama"], f"NIP. {d['nip']}"]
    for i, ln in enumerate(lines):
        par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        _fmt(par)
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _runs(par, ln)
    sig.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def build_super_docx(d, paper="A4"):
    doc, sec, cw = _setup_doc(paper)
    if QR_PATH.exists():
        sec.footer_distance = Cm(0.6)
        sec.footer.is_linked_to_previous = False
        fp = sec.footer.paragraphs[0]
        _fmt(fp)
        fp.paragraph_format.left_indent = Cm(-1.5)
        fp.add_run().add_picture(str(QR_PATH), width=Cm(2.0))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(p.add_run("SURAT PERNYATAAN"), 13, bold=True)
    _fmt(p, 24)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    _runs(p, "Yang bertanda tangan dibawah ini:")
    _fmt(p, 8)
    _kv_table(doc, [("Nama", d["nama"]), ("NIP", d["nip"]), ("Pangkat/Golongan", d["pangkat_gol"]),
                    ("Jabatan", d["jabatan"]), ("Unit Kerja", "Badan Pusat Statistik Kabupaten Bangka Selatan")],
              [4.2, 0.4, cw - 4.6])
    sp = doc.add_paragraph()
    _fmt(sp, 10)
    for txt in (d["p1"], d["p2"]):
        par = doc.add_paragraph()
        _fmt(par, 10, justify=True)
        _runs(par, txt)
    for _ in range(3):
        _fmt(doc.add_paragraph(), 6)
    sig = doc.add_table(rows=1, cols=2)
    sig.autofit = False
    sig.rows[0].cells[0].width = Cm(cw - 6.5)
    sig.rows[0].cells[1].width = Cm(6.5)
    cell = sig.rows[0].cells[1]
    lines = [f"Bangka Selatan, {d['tanggal']}", "Pelaksana Perjalanan Dinas", "", "", "",
             sign_name(d["nama"]), f"NIP. {d['nip']}"]
    for i, ln in enumerate(lines):
        par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        _fmt(par)
        _runs(par, ln)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ----------------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------------
def _rl(text):
    return ITALIC_RE.sub(r"<i>\1</i>", escape(text))


def _styles():
    base = ParagraphStyle("base", fontName="Times-Roman", fontSize=12, leading=15, alignment=TA_LEFT,
                          allowWidows=0, allowOrphans=0)
    return dict(
        base=base,
        title=ParagraphStyle("t", parent=base, fontName="Times-Bold", fontSize=13, alignment=TA_CENTER, spaceAfter=20),
        just=ParagraphStyle("j", parent=base, alignment=TA_JUSTIFY, spaceAfter=3),
        just10=ParagraphStyle("j10", parent=base, alignment=TA_JUSTIFY, spaceAfter=10),
        center=ParagraphStyle("c", parent=base, alignment=TA_CENTER),
    )


def _kv_flow(rows, widths):
    S = _styles()["base"]
    tbl = Table([[Paragraph(_rl(a), S), Paragraph(":", S), Paragraph(_rl(b), S)] for a, b in rows],
                colWidths=[w * cm for w in widths], hAlign="LEFT")
    tbl.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                             ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 0),
                             ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
    return tbl


def _sig_table(right, left_w, right_w):
    t = Table([["", right]], colWidths=[left_w * cm, right_w * cm], hAlign="LEFT")
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                           ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    return KeepTogether(t)


def _pdf_doc(buf, paper, title):
    w, h = PAPER[paper]
    return SimpleDocTemplate(buf, pagesize=(w * cm, h * cm), leftMargin=MARGIN_LR * cm,
                             rightMargin=MARGIN_LR * cm, topMargin=2 * cm, bottomMargin=2 * cm,
                             title=title), w - 2 * MARGIN_LR


def build_pdf(d, paper="A4"):
    buf = io.BytesIO()
    doc, cw = _pdf_doc(buf, paper, "Laporan Perjalanan Translok")
    S = _styles()
    story = []
    if d["kop"] and KOP_PATH.exists():
        iw, ih = Image.open(KOP_PATH).size
        story += [RLImage(str(KOP_PATH), width=KOP_WIDTH_CM * cm, height=KOP_WIDTH_CM * cm * ih / iw, hAlign="LEFT"),
                  Spacer(1, 14)]
    story.append(Paragraph("<u>Laporan Perjalanan Translok</u>", S["title"]))
    story += [_kv_flow([("Kepada", d["kepada"]), ("Perihal", d["perihal"]), ("Tujuan", d["tujuan"]),
                        ("Tanggal Pelaksanaan", d["tanggal"])], [4.2, 0.4, cw - 4.6]),
              Spacer(1, 14), Paragraph("Hasil Pelaksanaan", S["base"]), Spacer(1, 2)]
    n = 0
    for txt in [p for p in d["points"] if p.strip()] + ["Dokumentasi:"]:
        n += 1
        story.append(Paragraph(f"{n}. {_rl(txt)}", S["just"]))
    for img_bytes, w_cm, h_cm in d["photos"]:
        story.append(RLImage(io.BytesIO(img_bytes), width=w_cm * cm, height=h_cm * cm, hAlign="LEFT"))
        story.append(Spacer(1, 6))
    story.append(Spacer(1, 14))
    right = [Paragraph(f"Toboali, {d['tanggal']}", S["center"]), Paragraph("Pelaksana SPD", S["center"]),
             Spacer(1, 60), Paragraph(escape(d["nama"]), S["center"]), Paragraph(f"NIP. {d['nip']}", S["center"])]
    story.append(_sig_table(right, cw - 7.5, 7.5))
    doc.build(story)
    return buf.getvalue()


def build_super_pdf(d, paper="A4"):
    buf = io.BytesIO()
    doc, cw = _pdf_doc(buf, paper, "Surat Pernyataan")
    S = _styles()
    story = [Paragraph("<b>SURAT PERNYATAAN</b>", ParagraphStyle("tt", parent=S["title"], spaceAfter=24)),
             Paragraph("Yang bertanda tangan dibawah ini:", S["just10"]),
             _kv_flow([("Nama", d["nama"]), ("NIP", d["nip"]), ("Pangkat/Golongan", d["pangkat_gol"]),
                       ("Jabatan", d["jabatan"]), ("Unit Kerja", "Badan Pusat Statistik Kabupaten Bangka Selatan")],
                      [4.2, 0.4, cw - 4.6]),
             Spacer(1, 16), Paragraph(_rl(d["p1"]), S["just10"]), Paragraph(_rl(d["p2"]), S["just10"]),
             Spacer(1, 40)]
    right = [Paragraph(f"Bangka Selatan, {d['tanggal']}", S["base"]),
             Paragraph("Pelaksana Perjalanan Dinas", S["base"]), Spacer(1, 50),
             Paragraph(escape(sign_name(d["nama"])), S["base"]), Paragraph(f"NIP. {d['nip']}", S["base"])]
    story.append(_sig_table(right, cw - 6.5, 6.5))

    def qr(canvas, _doc):
        if QR_PATH.exists():
            canvas.drawImage(str(QR_PATH), 1 * cm, 0.8 * cm, 2 * cm, 2 * cm, mask="auto")

    doc.build(story, onFirstPage=qr, onLaterPages=qr)
    return buf.getvalue()


def render_preview(pdf_bytes):
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    pdf = fitz.open(stream=pdf_bytes, filetype="pdf")
    return [p.get_pixmap(dpi=100).tobytes("png") for p in pdf]


def zip_files(files):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files:
            z.writestr(name, data)
    return buf.getvalue()


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def safe(s):
    return re.sub(r"[^\w\-]+", "_", str(s)).strip("_")[:40]


# ----------------------------------------------------------------------------
# VALIDASI & EXCEL
# ----------------------------------------------------------------------------
def nk(s):
    return re.sub(r"[\s\.\,]+", "", str(s)).casefold()


def find_opt(value, options):
    m = {nk(o): o for o in options}
    return m.get(nk(value))


def display_desa(desa):
    if desa in DESA_DB:
        return pretty_desa(desa)
    return desa if re.match(r"(?i)(desa|kelurahan)\b", desa) else f"Desa {desa}"


def placeholders_left(text):
    return set(re.findall(r"\{(\w+)\}", text))


def validate_rec(rec):
    errs = []
    if rec["desa"] in DESA_DB and rec["kec"] in KECAMATAN and DESA_DB[rec["desa"]] != KECAMATAN[rec["kec"]]:
        errs.append(f"{display_desa(rec['desa'])} bukan bagian dari Kecamatan {KEC_DOC.get(rec['kec'], rec['kec'])}.")
    unit = "SLS" if rec["kegiatan"] == "Listing" else "sampel"
    if not rec["sls"].strip():
        errs.append(f"Nama {unit} belum diisi.")
    if not rec["pcl"].strip():
        errs.append("Nama PCL belum diisi.")
    need = 3 if rec["kegiatan"] == "Listing" else 1
    if len(rec["responden"]) < need:
        errs.append(f"Nama responden minimal {need} orang (terisi {len(rec['responden'])}).")
    if rec["tanggal"].year != rec["tahun"]:
        errs.append("Tanggal pelaksanaan tidak sesuai tahun yang dipilih.")
    return errs


def make_template(extra_vars):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    headers = DEFAULT_COLS + list(extra_vars)
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78" if h in DEFAULT_COLS else "C55A11")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = 48 if h == "Survei" else 24
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"

    lst = wb.create_sheet("Daftar")
    lists = {"Kegiatan": KEGIATAN, "Survei": SURVEI, "Tahun": TAHUN, "Subround": SUBROUND,
             "Jenis Ubinan": JENIS_UBINAN, "Triwulan": [1, 2, 3, 4], "Kecamatan": list(KECAMATAN),
             "Desa": [pretty_desa(k) for k in DESA_DB], "Nama Pelaksana": [n for n, _ in PELAKSANA]}
    for ci, (name, vals) in enumerate(lists.items(), 1):
        lst.cell(row=1, column=ci, value=name).font = Font(bold=True)
        for ri, val in enumerate(vals, 2):
            lst.cell(row=ri, column=ci, value=val)
        lst.column_dimensions[get_column_letter(ci)].width = 30
        if name in headers and not (name == "Desa" and "Kecamatan" in headers):
            col = get_column_letter(headers.index(name) + 1)
            L = get_column_letter(ci)
            # showErrorMessage=False -> boleh mengetik isian sendiri di luar dropdown
            dv = DataValidation(type="list", formula1=f"=Daftar!${L}$2:${L}${len(vals) + 1}",
                                allow_blank=True, showErrorMessage=False)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}500")

    # Dropdown Desa bergantung pada Kecamatan (dependent dropdown via named range + INDIRECT)
    if "Kecamatan" in headers and "Desa" in headers:
        from openpyxl.workbook.defined_name import DefinedName
        base_col = len(lists) + 2
        for j, (label, kdb) in enumerate(KECAMATAN.items()):
            ci = base_col + j
            desa_k = [pretty_desa(k) for k, v in DESA_DB.items() if v == kdb]
            lst.cell(row=1, column=ci, value=f"Desa - {label}").font = Font(bold=True)
            for ri, val in enumerate(desa_k, 2):
                lst.cell(row=ri, column=ci, value=val)
            lst.column_dimensions[get_column_letter(ci)].width = 30
            nm = "Desa_" + label.replace(" ", "_")
            L = get_column_letter(ci)
            wb.defined_names[nm] = DefinedName(nm, attr_text=f"Daftar!${L}$2:${L}${len(desa_k) + 1}")
        kc = get_column_letter(headers.index("Kecamatan") + 1)
        dc = get_column_letter(headers.index("Desa") + 1)
        all_L = get_column_letter(list(lists).index("Desa") + 1)
        all_rng = f"Daftar!${all_L}$2:${all_L}${len(lists['Desa']) + 1}"
        dv = DataValidation(type="list", allow_blank=True, showErrorMessage=False,
                            formula1=f'=IFERROR(INDIRECT("Desa_"&SUBSTITUTE({kc}2," ","_")),{all_rng})')
        ws.add_data_validation(dv)
        dv.add(f"{dc}2:{dc}500")

    pt = wb.create_sheet("Petunjuk")
    notes = [
        "PETUNJUK PENGISIAN",
        "1. Isi sheet 'Data' mulai baris 2. Satu baris = satu dokumen laporan.",
        "2. Kolom biru = variabel default, kolom oranye = variabel tambahan (dipakai di teks sebagai {nama_kolom}).",
        "3. Dropdown hanya BANTUAN. Anda boleh mengetik isian sendiri (survei, kegiatan, kecamatan, desa, pelaksana, dll.) di luar daftar.",
        "4. Subround & Jenis Ubinan diisi bila Survei = Survei Ubinan. Triwulan diisi untuk survei berakhiran 'Triwulan B'.",
        "5. Pilih Kecamatan dulu; dropdown Desa otomatis menyesuaikan kecamatan terpilih. Jika diketik sendiri, tidak dicek.",
        "6. Nama Responden dipisah koma. Listing minimal 3 responden.",
        "7. Tanggal Pelaksanaan berformat tanggal (mis. 24/04/2026) dan harus sesuai Tahun.",
        "8. Nama Pelaksana di luar daftar: isi juga kolom 'NIP Pelaksana'. Untuk nama dalam daftar, NIP terisi otomatis.",
        "9. Nama File Foto (opsional): nama file foto yang diupload di aplikasi, pisahkan koma bila lebih dari satu.",
        "",
        "Contoh: Listing | Survei Ubinan | 2026 | Subround 2 (SR 2) | Palawija | | Tukak Sadai | Bukit Terap |",
        "RT 007 Dusun Terap | Nupriyadi | Amri, Sukri, Aripin | 24/04/2026 | Stenislaus Angga Aprianto, S.Tr.Stat | | foto1.jpg",
    ]
    for i, t in enumerate(notes, 1):
        pt.cell(row=i, column=1, value=t)
    pt["A1"].font = Font(bold=True)
    pt.column_dimensions["A"].width = 135
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def cell_str(x):
    if x is None or x is pd.NaT:
        return ""
    try:
        if pd.isna(x):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(x, float) and x.is_integer():
        return str(int(x))
    return str(x).strip()


def parse_excel(file):
    xl = pd.ExcelFile(file)
    sheet = "Data" if "Data" in xl.sheet_names else xl.sheet_names[0]
    df = xl.parse(sheet, dtype=object)
    df.columns = [str(c).strip() for c in df.columns]
    extra_cols = [c for c in df.columns if c not in DEFAULT_COLS and not c.startswith("Unnamed")]
    cores = {k.split(" ", 1)[1]: k for k in DESA_DB}
    peg_names = [n for n, _ in PELAKSANA]
    recs, errors = [], []
    for idx, row in df.iterrows():
        cols = [c for c in DEFAULT_COLS + extra_cols if c in df.columns]
        vals = {c: row.get(c) for c in cols}
        if all(cell_str(x) == "" for x in vals.values()):
            continue
        rn = idx + 2

        def g(c):
            return cell_str(vals.get(c))

        errs = []
        keg = find_opt(g("Kegiatan"), KEGIATAN) or g("Kegiatan")
        if not keg:
            errs.append("Kegiatan kosong")
        sur = find_opt(g("Survei"), SURVEI) or g("Survei")
        if not sur:
            errs.append("Survei kosong")
        try:
            thn = int(float(g("Tahun")))
        except ValueError:
            thn = None
            errs.append("Tahun tidak valid")
        sub = jen = tri = None
        if sur == "Survei Ubinan":
            s = g("Subround")
            sub = SUBROUND[int(s) - 1] if s in ("1", "2", "3") else (find_opt(s, SUBROUND) or s)
            jen = find_opt(g("Jenis Ubinan"), JENIS_UBINAN) or g("Jenis Ubinan")
            if not sub:
                errs.append("Subround kosong")
            if not jen:
                errs.append("Jenis Ubinan kosong")
        if sur and has_B(sur):
            tri = g("Triwulan")
            if not tri:
                errs.append("Triwulan harus diisi")
        kec_in = g("Kecamatan")
        kec = find_opt(kec_in, list(KECAMATAN)) or ("Lepar" if nk(kec_in) == nk("Lepar Pongok") else kec_in)
        if not kec:
            errs.append("Kecamatan kosong")
        d_in = re.sub(r"^(desa|kelurahan|kel\.?)\s+", "", g("Desa"), flags=re.I).strip()
        desa = ""
        if not d_in:
            errs.append("Desa kosong")
        else:
            m = find_opt(d_in, list(cores))
            desa = cores[m] if m else g("Desa")
        raw_t = vals.get("Tanggal Pelaksanaan")
        tgl = None
        if isinstance(raw_t, (datetime, date)) and not pd.isna(raw_t):
            tgl = pd.Timestamp(raw_t).date()
        else:
            ts = pd.to_datetime(cell_str(raw_t), dayfirst=True, errors="coerce")
            tgl = None if pd.isna(ts) else ts.date()
        if tgl is None:
            errs.append("Tanggal tidak valid")
        pel_in = g("Nama Pelaksana")
        pn = find_opt(pel_in, peg_names)
        pel_nama = pel_nip = ""
        if pn:
            pel_nama, pel_nip = PELAKSANA[peg_names.index(pn)]
        elif pel_in:
            pel_nama, pel_nip = pel_in, g("NIP Pelaksana")
            if not pel_nip:
                errs.append("NIP Pelaksana harus diisi untuk nama di luar daftar")
        else:
            errs.append("Nama Pelaksana kosong")
        if not errs:
            rec = dict(kegiatan=keg, survei=sur, tahun=thn, subround=sub, jenis=jen, triwulan=tri, kec=kec,
                       desa=desa, sls=g("Nama SLS/Sampel"), pcl=g("Nama PCL"),
                       responden=split_names(g("Nama Responden")), tanggal=tgl,
                       pelaksana_nama=pel_nama, pelaksana_nip=pel_nip,
                       extra={var_key(c): g(c) for c in extra_cols},
                       foto=[x.strip() for x in re.split(r"[,;\n]+", g("Nama File Foto")) if x.strip()])
            errs += validate_rec(rec)
        if errs:
            errors.append((rn, "; ".join(errs)))
        else:
            rec["row"] = rn
            recs.append(rec)
    return recs, errors, extra_cols


# ----------------------------------------------------------------------------
# STATE & KOMPONEN BERSAMA
# ----------------------------------------------------------------------------
def init_state():
    ss = st.session_state
    ss.setdefault("styles", {"Default": copy.deepcopy(DEFAULT_STYLE)})
    ss.setdefault("styles_ver", 0)
    ss.setdefault("applied_style", "Default")
    ss.setdefault("extra_vars", [])
    ss.setdefault("xl_edver", 0)


def reset_inputs(prefixes, counter=None, extra_del=()):
    """Reset HANYA isian data. Gaya teks, templat, dan poin uraian tidak dihapus."""
    ss = st.session_state
    for k in list(ss.keys()):
        if k.startswith(tuple(prefixes)) or k in extra_del:
            del ss[k]
    if counter:
        ss[counter] = ss.get(counter, 0) + 1


LAP_INPUTS = ("lap_kegiatan", "lap_tahun", "lap_survei", "lap_subround", "lap_jenis", "lap_triwulan",
              "lap_kec", "lap_desa_", "lap_sls", "lap_pcl", "lap_resp", "lap_tgl_", "lap_pel",
              "lap_files_", "lap_mode", "lap_var_")
SP_INPUTS = ("sp_peg", "sp_nomor", "sp_tgl")


def restore(keys):
    """Streamlit membuang state widget yang tidak ditampilkan; pulihkan dari salinan agar isian tidak hilang."""
    ss = st.session_state
    for k in keys:
        if k not in ss and ("_keep_" + k) in ss:
            ss[k] = ss["_keep_" + k]


def save(keys):
    ss = st.session_state
    for k in keys:
        if k in ss:
            ss["_keep_" + k] = ss[k]


def style_for(name, keg):
    s = st.session_state["styles"].get(name) or st.session_state["styles"]["Default"]
    key = keg if keg in KEGIATAN else "Pendataan"
    return s.get(key) or DEFAULT_STYLE[key]


def all_var_names():
    return BASE_VARS + [var_key(c) for c in st.session_state["extra_vars"]]


def load_tpl_into_editor(prefix, name, keg):
    t = style_for(name, keg)
    for k in ("perihal", "p1", "p2", "p3"):
        st.session_state[f"{prefix}_tpl_{k}"] = t[k]


def cb_save_style(prefix, keg):
    ss = st.session_state
    new = (ss.get(f"{prefix}_newstyle") or "").strip()
    name = new or ss.get(f"{prefix}_style_sel", "Default")
    if name == "Default":
        ss[f"{prefix}_style_msg"] = ("error", "Isi 'Nama gaya baru' dulu (gaya 'Default' tidak bisa ditimpa).")
        return
    ss["styles"].setdefault(name, copy.deepcopy(DEFAULT_STYLE))
    ss["styles"][name][keg] = {k: ss[f"{prefix}_tpl_{k}"] for k in ("perihal", "p1", "p2", "p3")}
    ss["styles_ver"] += 1
    ss[f"{prefix}_style_sel"] = name
    ss[f"{prefix}_tplsig"] = (keg, name, ss["styles_ver"])
    ss[f"{prefix}_newstyle"] = ""
    ss[f"{prefix}_style_msg"] = ("success", f"Gaya '{name}' untuk kegiatan {keg} tersimpan.")


def cb_apply_style(prefix, keg):
    ss = st.session_state
    name = ss.get(f"{prefix}_style_sel", "Default")
    ss["applied_style"] = name
    load_tpl_into_editor(prefix, name, keg)
    ss[f"{prefix}_style_msg"] = ("success", f"Gaya '{name}' diterapkan ke semua dokumen.")


def style_editor(prefix, keg, save_label, apply_fmt):
    ss = st.session_state
    styles = ss["styles"]
    keep_keys = [f"{prefix}_style_sel"] + [f"{prefix}_tpl_{k}" for k in ("perihal", "p1", "p2", "p3")]
    restore(keep_keys)
    sel = st.selectbox("Pilih Gaya Teks (X)", list(styles), key=f"{prefix}_style_sel")
    sig_t = (keg, sel, ss["styles_ver"])
    if ss.get(f"{prefix}_tplsig") != sig_t or f"{prefix}_tpl_p1" not in ss:
        ss[f"{prefix}_tplsig"] = sig_t
        load_tpl_into_editor(prefix, sel, keg)
    st.caption("Ubah templat kalimat. Variabel dalam kurung kurawal otomatis diganti: " +
               ", ".join("{" + x + "}" for x in all_var_names()))
    st.text_area("Templat Perihal", key=f"{prefix}_tpl_perihal", height=70)
    st.text_area("Templat Poin 1", key=f"{prefix}_tpl_p1", height=90)
    st.text_area("Templat Poin 2", key=f"{prefix}_tpl_p2", height=100)
    st.text_area("Templat Poin 3", key=f"{prefix}_tpl_p3", height=90)
    save(keep_keys)
    st.text_input("Nama gaya baru", key=f"{prefix}_newstyle", placeholder="mis. Gaya Singkat")
    b1, b2 = st.columns(2)
    b1.button(save_label, on_click=cb_save_style, args=(prefix, keg), key=f"btn_save_style_{prefix}",
              use_container_width=True)
    b2.button(apply_fmt.format(X=sel), on_click=cb_apply_style, args=(prefix, keg),
              key=f"btn_apply_style_{prefix}", use_container_width=True)
    msg = ss.get(f"{prefix}_style_msg")
    if msg:
        (st.success if msg[0] == "success" else st.error)(msg[1])
    st.caption(f"Gaya yang diterapkan ke semua dokumen: **{ss['applied_style']}**")
    st.download_button("⬇️ Unduh semua gaya (JSON)", json.dumps(styles, ensure_ascii=False, indent=1),
                       file_name="gaya_teks_laporan.json", mime="application/json", key=f"dl_styles_{prefix}")
    up = st.file_uploader("Muat gaya dari file JSON", type=["json"], key=f"{prefix}_styleup")
    if up is not None and ss.get("styleup_last") != (up.name, up.size):
        try:
            data = json.loads(up.getvalue().decode("utf-8"))
            for k, val in data.items():
                if k != "Default" and isinstance(val, dict):
                    ss["styles"][k] = {kg: {x: str(val.get(kg, DEFAULT_STYLE[kg])[x]) for x in
                                            ("perihal", "p1", "p2", "p3")} for kg in KEGIATAN}
            ss["styleup_last"] = (up.name, up.size)
            ss["styles_ver"] += 1
            st.rerun()
        except Exception:
            st.error("File JSON gaya tidak valid.")
    return sel


# ---- variabel tambahan ----
def cb_add_var(prefix):
    ss = st.session_state
    name = (ss.get(f"{prefix}_newvar") or "").strip()
    vk = var_key(name)
    existing = [var_key(c) for c in ss["extra_vars"]]
    if not vk:
        ss["var_msg"] = ("error", "Nama variabel belum diisi.")
    elif vk in BASE_VARS or vk in existing or name in DEFAULT_COLS:
        ss["var_msg"] = ("error", f"Variabel '{name}' sudah ada / bentrok dengan variabel default.")
    else:
        ss["extra_vars"].append(name)
        ss[f"{prefix}_newvar"] = ""
        ss["var_msg"] = ("success", f"Variabel ditambahkan. Pakai di teks sebagai {{{vk}}}.")


def cb_del_var(i):
    st.session_state["extra_vars"].pop(i)


def var_manager(prefix):
    ss = st.session_state
    st.text_input("Nama variabel tambahan", key=f"{prefix}_newvar", placeholder="mis. Nama Koordinator")
    st.button("➕ Tambahkan Variabel Selain Variabel Default", on_click=cb_add_var, args=(prefix,),
              key=f"btn_addvar_{prefix}")
    msg = ss.pop("var_msg", None)
    if msg:
        (st.success if msg[0] == "success" else st.error)(msg[1])
    for i, c in enumerate(ss["extra_vars"]):
        c1, c2 = st.columns([6, 1])
        c1.caption(f"• {c} → {{{var_key(c)}}}")
        c2.button("🗑️", key=f"btn_delvar_{prefix}_{i}", on_click=cb_del_var, args=(i,))


# ---- poin tambahan ----
def cb_insert_var(prefix, n):
    ss = st.session_state
    v = ss.get(f"{prefix}_insvar_{n}")
    if v:
        ss[f"{prefix}_newpoint"] = (ss.get(f"{prefix}_newpoint") or "") + "{" + v + "}"


def cb_add_point(prefix):
    ss = st.session_state
    txt = (ss.get(f"{prefix}_newpoint") or "").strip()
    if txt:
        ss.setdefault(f"{prefix}_extras", []).append(
            {"text": txt, "after": POS_OPTIONS.index(ss.get(f"{prefix}_newpos", POS_OPTIONS[0]))})
        ss[f"{prefix}_newpoint"] = ""


def cb_del_point(prefix, i):
    st.session_state[f"{prefix}_extras"].pop(i)


def points_editor(prefix):
    ss = st.session_state
    n = len(ss["extra_vars"])
    st.markdown("**Tambah poin uraian**")
    st.text_area("Uraian poin tambahan (boleh memakai variabel)", key=f"{prefix}_newpoint", height=80,
                 placeholder="mis. Pada keluarga {nama_responden} terdapat {nama_anggota} dengan komoditas utamanya {komoditas}")
    c1, c2 = st.columns([3, 2])
    c1.selectbox("Variabel yang disisipkan", all_var_names(), key=f"{prefix}_insvar_{n}")
    c2.markdown("<div style='height:1.9em'></div>", unsafe_allow_html=True)
    c2.button("↳ Sisipkan variabel", on_click=cb_insert_var, args=(prefix, n), key=f"btn_insvar_{prefix}")
    st.selectbox("Posisi poin", POS_OPTIONS, key=f"{prefix}_newpos",
                 help="Sisipkan setelah poin x, atau biarkan pada urutan semestinya.")
    st.button("➕ Tambah poin", on_click=cb_add_point, args=(prefix,), key=f"btn_addpoint_{prefix}")
    for i, e in enumerate(ss.get(f"{prefix}_extras", [])):
        a, b = st.columns([6, 1])
        pos = "urutan semestinya" if e["after"] == 0 else f"setelah poin {e['after']}"
        a.caption(f"• ({pos}) {e['text']}")
        b.button("🗑️", key=f"btn_delpoint_{prefix}_{i}", on_click=cb_del_point, args=(prefix, i))


# ----------------------------------------------------------------------------
# DOWNLOAD
# ----------------------------------------------------------------------------
def download_block(docs, paper, key, base_name):
    fmt = st.radio("Pilih format download", ["PDF", "Word (.docx)", "PDF + Word"], horizontal=True, key=key + "_fmt")
    files = []
    for i, (label, d) in enumerate(docs, 1):
        nm = base_name if len(docs) == 1 else f"{base_name}_{i:02d}"
        if fmt in ("PDF", "PDF + Word"):
            files.append((nm + ".pdf", d["_pdf"]))
        if fmt in ("Word (.docx)", "PDF + Word"):
            files.append((nm + ".docx", build_docx(d, paper)))
    if len(files) == 1:
        n, data = files[0]
        st.download_button("⬇️ Download " + n.rsplit(".", 1)[1].upper(), data, file_name=n,
                           mime="application/pdf" if n.endswith(".pdf") else DOCX_MIME, type="primary", key=key + "_dl")
    else:
        st.download_button(f"⬇️ Download ZIP ({len(files)} file)", zip_files(files), file_name=base_name + ".zip",
                           mime="application/zip", type="primary", key=key + "_dl")


# ----------------------------------------------------------------------------
# HALAMAN: LAPORAN MANUAL
# ----------------------------------------------------------------------------
def cb_reset_texts():
    st.session_state["lap_sig"] = None


def page_laporan(paper, use_kop, photo_w, photo_h):
    ss = st.session_state
    st.title("📝 Laporan Perjalanan Translok")
    col_form, col_prev = st.columns([1, 1.15], gap="large")

    with col_form:
        h1, h2 = st.columns([2, 1])
        h1.subheader("1. Isi Data")
        h2.button("🔄 Reset isian data", on_click=reset_inputs, args=(LAP_INPUTS, "upl_n_lap_"),
                  key="btn_reset_lap", use_container_width=True,
                  help="Hanya mengosongkan isian data. Gaya teks, templat, dan poin uraian tetap ada.")
        c1, c2 = st.columns(2)
        kegiatan = c1.selectbox("Pilih Kegiatan", KEGIATAN, key="lap_kegiatan")
        tahun = c2.selectbox("Pilih Tahun", TAHUN, key="lap_tahun")
        survei = st.selectbox("Pilih Survei", SURVEI, key="lap_survei")
        subround = jenis = triwulan = None
        if survei == "Survei Ubinan":
            c1, c2 = st.columns(2)
            subround = c1.selectbox("Pilih Subround", SUBROUND, key="lap_subround")
            jenis = c2.selectbox("Jenis Ubinan", JENIS_UBINAN, key="lap_jenis")
        if has_B(survei):
            triwulan = st.selectbox("Pilih Triwulan (nilai B)", [1, 2, 3, 4], key="lap_triwulan")
        kec = st.selectbox("Tujuan: Kecamatan", list(KECAMATAN), key="lap_kec")
        desa_opts = [k for k, v in DESA_DB.items() if v == KECAMATAN[kec]]
        desa = st.selectbox("Pilih Desa (otomatis sesuai kecamatan)", desa_opts, format_func=pretty_desa,
                            key=f"lap_desa_{kec}")
        unit = "SLS" if kegiatan == "Listing" else "sampel"
        sls = st.text_input(f"Ketikkan nama {unit}", key="lap_sls",
                            placeholder="RT 007 Dusun Terap" if kegiatan == "Listing" else "BPP Air Gegas")
        pcl = st.text_input("Ketik Nama PCL (1 PCL)", key="lap_pcl")
        resp_label = ("Ketik Nama Responden (min. 3 responden)" if kegiatan == "Listing"
                      else "Ketik Nama Responden / Penanggungjawab Usaha")
        resp_txt = st.text_area(resp_label, help="Pisahkan dengan koma atau baris baru.", height=90, key="lap_resp")
        tgl_default = date.today() if date.today().year == tahun else date(tahun, 1, 1)
        tgl = st.date_input("Pilih Tanggal", value=tgl_default, min_value=date(tahun, 1, 1),
                            max_value=date(tahun, 12, 31), key=f"lap_tgl_{tahun}", format="DD/MM/YYYY")
        pel = st.selectbox("Nama Pelaksana", range(len(PELAKSANA)), format_func=lambda i: PELAKSANA[i][0], key="lap_pel")

        with st.expander("➕ Variabel tambahan (untuk dipakai di uraian poin)"):
            var_manager("lap")
        extra = {}
        if ss["extra_vars"]:
            st.markdown("**Isi variabel tambahan**")
            for c in ss["extra_vars"]:
                extra[var_key(c)] = st.text_input(c, key=f"lap_var_{var_key(c)}")

        files = st.file_uploader("Upload Dokumentasi (JPG/JPEG, maks 5 MB per foto)", type=["jpg", "jpeg"],
                                 accept_multiple_files=True, key=f"lap_files_{ss.get('upl_n_lap_', 0)}")
        photos, perr = [], []
        for f in files or []:
            if f.size > 5 * 1024 * 1024:
                perr.append(f"Foto '{f.name}' melebihi 5 MB.")
                continue
            try:
                photos.append(process_photo_bytes(f.getvalue()))
            except Exception:
                perr.append(f"Foto '{f.name}' tidak bisa dibaca.")
        photos = scale_photos(photos, photo_w, photo_h, paper)
        mode_single = True
        if len(photos) > 1:
            mode = st.radio("Gambar lebih dari 1, apakah akan dimasukkan ke dalam 1 dokumen yang sama? "
                            "atau masukkan gambar 1 dokumen untuk satu gambar saja!",
                            ["Masukkan semua gambar ke dalam 1 dokumen yang sama",
                             "1 dokumen untuk 1 gambar (dokumen terpisah)"], key="lap_mode")
            mode_single = mode.startswith("Masukkan")

    rec = dict(kegiatan=kegiatan, survei=survei, tahun=tahun, subround=subround, jenis=jenis, triwulan=triwulan,
               kec=kec, desa=desa, sls=sls, pcl=pcl, responden=split_names(resp_txt), tanggal=tgl,
               pelaksana_nama=PELAKSANA[pel][0], pelaksana_nip=PELAKSANA[pel][1], extra=extra)
    errors = validate_rec(rec) + perr
    for e in errors:
        col_form.error(e)

    with col_prev:
        st.subheader("2. Preview & Edit")
        if errors:
            st.info("Lengkapi/perbaiki data di sebelah kiri untuk menampilkan preview.")
            return
        v = make_vars(rec)

        with st.expander("🎨 Gaya Penulisan Teks Laporan"):
            style_editor("lap", kegiatan, "💾 Simpan Gaya Penulisan Teks Laporan",
                         "✅ Terapkan Gaya Teks {X} Pada Seluruh Dokumen")

        auto = {k: render_tpl(ss[f"lap_tpl_{k}"], v) for k in ("perihal", "p1", "p2", "p3")}
        sig = (tuple(sorted(v.items())), tuple(auto.values()))
        f_keys = [f"lap_f_{k}" for k in ("perihal", "p1", "p2", "p3")]
        restore(f_keys)
        if ss.get("lap_sig") != sig or any(k not in ss for k in f_keys):
            ss["lap_sig"] = sig
            for k, t in auto.items():
                ss[f"lap_f_{k}"] = t

        with st.expander("✏️ Edit teks & tambah poin uraian"):
            st.text_area("Perihal", key="lap_f_perihal", height=80)
            st.text_area("Poin 1", key="lap_f_p1", height=100)
            st.text_area("Poin 2", key="lap_f_p2", height=110)
            st.text_area("Poin 3", key="lap_f_p3", height=100)
            save(f_keys)
            st.button("↩️ Kembalikan ke teks otomatis", on_click=cb_reset_texts, key="btn_reset_texts")
            st.divider()
            points_editor("lap")

        points = assemble_points([ss["lap_f_p1"], ss["lap_f_p2"], ss["lap_f_p3"]], ss.get("lap_extras", []), v)
        left = placeholders_left(ss["lap_f_perihal"] + " " + " ".join(points))
        if left:
            st.warning("Variabel belum terisi/terdefinisi: " + ", ".join("{" + x + "}" for x in sorted(left)))
        empty = [x for x in placeholders_left(" ".join(points)) if x in extra and not extra[x]]
        groups = [photos] if (mode_single or len(photos) <= 1) else [[p] for p in photos]
        docs = []
        for g in groups:
            d = doc_data(rec, v, ss["lap_f_perihal"], points, g, use_kop)
            d["_pdf"] = build_pdf(d, paper)
            docs.append((f"Dokumen {len(docs) + 1}", d))
        for label, d in docs:
            if len(docs) > 1:
                st.markdown(f"**{label}**")
            for png in render_preview(d["_pdf"]):
                st.image(png, use_container_width=True)
        st.caption("Preview dibuat dari versi PDF. Tata letak file Word bisa sedikit berbeda.")
        st.subheader("3. Download")
        download_block(docs, paper, "lap", f"Laporan_Translok_{tgl.strftime('%Y-%m-%d')}")


# ----------------------------------------------------------------------------
# HALAMAN: EXCEL
# ----------------------------------------------------------------------------
def gen_templates(keg):
    ss = st.session_state
    key = keg if keg in KEGIATAN else "Pendataan"
    if key == ss.get("xl_keg_edit") and all(f"xl_tpl_{k}" in ss for k in ("perihal", "p1", "p2", "p3")):
        return {k: ss[f"xl_tpl_{k}"] for k in ("perihal", "p1", "p2", "p3")}
    return style_for(ss["applied_style"], key)


def cb_add_doc_point(idx):
    ss = st.session_state
    item = ss["xl_docs"][idx]
    d = item["d"]
    n = len(d["points"])
    txt = (ss.get(f"xl_dnew_{idx}") or "").strip()
    if not txt:
        return
    pos = ss.get(f"xl_dpos_{idx}_{n}", 0)
    txt = render_tpl(txt, item["v"])
    if pos == 0:
        d["points"].append(txt)
    else:
        d["points"].insert(pos, txt)
    ss["xl_edver"] += 1
    ss["xl_zip"] = None
    ss[f"xl_dnew_{idx}"] = ""


def page_excel(paper, use_kop, photo_w, photo_h):
    ss = st.session_state
    st.title("📊 Generate Laporan dari Excel")

    st.markdown("#### 1. Siapkan template Excel")
    c1, c2 = st.columns([3, 2])
    with c1:
        st.caption("Kolom template sama dengan isian form manual. Tambahkan variabel lain bila perlu "
                   "(mis. 'Nama Koordinator' → {nama_koordinator}).")
        var_manager("xl")
    with c2:
        st.write("**Kolom template:**")
        st.write(", ".join(DEFAULT_COLS + ss["extra_vars"]))
        st.download_button("⬇️ Download Template Excel", make_template(ss["extra_vars"]),
                           file_name="Template_Laporan_Translok.xlsx", key="dl_tpl",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary")
    st.info("Dropdown di Excel hanya bantuan. Anda juga boleh mengetik isian sendiri di luar daftar "
            "(survei, kegiatan, kecamatan, desa, pelaksana, dll.); aplikasi tetap membacanya.")

    st.markdown("#### 2. Editor isian dokumen")
    keg_edit = st.radio("Kegiatan yang diedit", KEGIATAN, horizontal=True, key="xl_keg_edit")
    with st.expander("🎨 Gaya Penulisan Teks", expanded=True):
        style_editor("xl", keg_edit, "💾 Simpan Gaya Penulisan Ini", "✅ Terapkan Gaya Penulisan {X} ke semua dokumen")
        st.caption("Untuk kegiatan lain yang tidak sedang diedit dipakai gaya yang diterapkan. "
                   "Simpan gaya sebelum berganti kegiatan agar editan tidak hilang.")
    with st.expander("✏️ Edit & tambah poin uraian (berlaku untuk semua dokumen)", expanded=True):
        points_editor("xl")

    st.markdown("#### 3. Upload Excel & generate")
    n_upl = ss.get("upl_n_xl_", 0)
    xl = st.file_uploader("Upload file Excel yang sudah diisi", type=["xlsx"], key=f"xl_file_{n_upl}")
    fotos = st.file_uploader("Upload foto dokumentasi (opsional, sesuai kolom 'Nama File Foto')",
                             type=["jpg", "jpeg"], accept_multiple_files=True, key=f"xl_fotos_{n_upl}")
    st.button("🔄 Reset upload", on_click=reset_inputs, args=(("xl_file_", "xl_fotos_"), "upl_n_xl_",
                                                              ("xl_docs", "xl_zip")), key="btn_reset_xl",
              help="Hanya mengosongkan file upload & hasil generate. Gaya dan poin tetap ada.")
    if xl:
        try:
            recs, errors, _ = parse_excel(xl)
        except Exception as e:
            st.error(f"File Excel tidak bisa dibaca: {e}")
            recs, errors = [], []
        if recs or errors:
            st.success(f"{len(recs)} baris valid.")
        if errors:
            st.error(f"{len(errors)} baris bermasalah dan dilewati:")
            st.dataframe(pd.DataFrame(errors, columns=["Baris Excel", "Masalah"]), hide_index=True)
        fmap = {f.name.casefold(): f for f in fotos or []}
        if recs and st.button("🚀 Generate laporan dari Excel", type="primary", key="btn_gen_xl"):
            docs, warns, left = [], [], set()
            prog = st.progress(0.0)
            for i, rec in enumerate(recs, 1):
                v = make_vars(rec)
                tpl = gen_templates(rec["kegiatan"])
                texts = [render_tpl(tpl[k], v) for k in ("p1", "p2", "p3")]
                points = assemble_points(texts, ss.get("xl_extras", []), v)
                perihal = render_tpl(tpl["perihal"], v)
                ph = []
                for fn in rec["foto"]:
                    f = fmap.get(fn.casefold())
                    if f is None:
                        warns.append(f"Baris {rec['row']}: foto '{fn}' tidak ditemukan di upload.")
                    else:
                        try:
                            ph.append(process_photo_bytes(f.getvalue()))
                        except Exception:
                            warns.append(f"Baris {rec['row']}: foto '{fn}' tidak bisa dibaca.")
                left |= placeholders_left(perihal + " " + " ".join(points))
                d = doc_data(rec, v, perihal, points, scale_photos(ph, photo_w, photo_h, paper), use_kop)
                docs.append(dict(name=f"{i:02d}_{safe(v['tanggal'])}_{safe(v['kecamatan'])}_{safe(rec['pcl'])}",
                                 label=f"Baris {rec['row']} – {rec['pcl']} – {v['kecamatan']}", d=d, v=v))
                prog.progress(i / len(recs))
            if left:
                warns.append("Variabel belum terdefinisi/tidak ada kolomnya di Excel: " +
                             ", ".join("{" + x + "}" for x in sorted(left)))
            ss["xl_docs"], ss["xl_warns"], ss["xl_zip"] = docs, warns, None
            ss["xl_edver"] += 1

    docs = ss.get("xl_docs")
    if not docs:
        return
    for w in ss.get("xl_warns", []):
        st.warning(w)
    st.success(f"{len(docs)} dokumen siap. Pilih satu dokumen untuk dicek / diedit satu per satu.")
    idx = st.selectbox("Pilih dokumen", range(len(docs)), format_func=lambda i: f"{i + 1}. {docs[i]['label']}",
                       key="xl_pick")
    item = docs[idx]
    d = item["d"]
    ev = ss["xl_edver"]
    ce, cp = st.columns([1, 1.1], gap="large")
    with ce:
        with st.expander("✏️ Edit dokumen ini (hanya dokumen terpilih)", expanded=True):
            np_ = st.text_area("Perihal", d["perihal"], key=f"xled_{ev}_{idx}_perihal", height=80)
            pts = [st.text_area(f"Poin {j + 1}", p, key=f"xled_{ev}_{idx}_p{j}", height=110)
                   for j, p in enumerate(d["points"])]
            if np_ != d["perihal"] or pts != d["points"]:
                d["perihal"], d["points"] = np_, pts
                ss["xl_zip"] = None
            st.caption("Kosongkan teks sebuah poin untuk menghapusnya dari dokumen.")
            st.divider()
            st.text_area("Tambah poin untuk dokumen ini (boleh memakai variabel)", key=f"xl_dnew_{idx}", height=80)
            n = len(d["points"])
            opts = [POS_OPTIONS[0]] + [f"Sisipkan setelah poin {j}" for j in range(1, n + 1)]
            st.selectbox("Posisi poin", range(len(opts)), format_func=lambda i: opts[i], key=f"xl_dpos_{idx}_{n}")
            st.button("➕ Tambah poin ke dokumen ini", on_click=cb_add_doc_point, args=(idx,),
                      key=f"btn_adddoc_{idx}")
    with cp:
        st.markdown("**Preview**")
        for png in render_preview(build_pdf(d, paper)):
            st.image(png, use_container_width=True)

    st.markdown("#### 4. Download")
    fmt = st.radio("Format hasil", ["PDF", "Word (.docx)", "PDF + Word"], horizontal=True, key="xl_fmt")
    if st.button("📦 Siapkan file untuk download", key="btn_zip"):
        files = []
        with st.spinner("Membuat file..."):
            for it in docs:
                if fmt in ("PDF", "PDF + Word"):
                    files.append((it["name"] + ".pdf", build_pdf(it["d"], paper)))
                if fmt in ("Word (.docx)", "PDF + Word"):
                    files.append((it["name"] + ".docx", build_docx(it["d"], paper)))
        ss["xl_zip"] = zip_files(files)
    if ss.get("xl_zip"):
        st.download_button("⬇️ Download semua (ZIP)", ss["xl_zip"], file_name="Laporan_Translok_dari_Excel.zip",
                           mime="application/zip", type="primary", key="dl_xl_zip")
    else:
        st.caption("Klik 'Siapkan file untuk download' (diulang bila ada edit baru).")


# ----------------------------------------------------------------------------
# HALAMAN: SURAT PERNYATAAN
# ----------------------------------------------------------------------------
def cb_reset_super_text():
    st.session_state["sp_sig"] = None


def page_super(paper):
    ss = st.session_state
    st.title("🚗 Surat Pernyataan Tidak Menggunakan Kendaraan Dinas (Super)")
    col_form, col_prev = st.columns([1, 1.15], gap="large")
    with col_form:
        h1, h2 = st.columns([2, 1])
        h1.subheader("1. Isi Data")
        h2.button("🔄 Reset isian data", on_click=reset_inputs, args=(SP_INPUTS,), key="btn_reset_sp",
                  use_container_width=True)
        idx = st.selectbox("Nama Pegawai", range(len(PEGAWAI)), format_func=lambda i: PEGAWAI[i]["nama"],
                           key="sp_peg", index=11)
        peg = PEGAWAI[idx]
        pg = f"{peg['pangkat']}/ {peg['gol']}" if peg["pangkat"] else peg["gol"]
        st.text_input("NIP", peg["nip"], disabled=True, key=f"sp_nip_{idx}")
        st.text_input("Pangkat/Golongan", pg, disabled=True, key=f"sp_pg_{idx}")
        st.text_input("Jabatan", peg["jabatan"], disabled=True, key=f"sp_jab_{idx}")
        nomor = st.text_input("Nomor Surat Tugas", key="sp_nomor", placeholder="B-1430/01905-ST/KP.650/2026")
        tgl = st.date_input("Pilih Tanggal (pelaksanaan & tanda tangan)", value=date.today(),
                            min_value=date(2026, 1, 1), max_value=date(2030, 12, 31), key="sp_tgl", format="DD/MM/YYYY")
    with col_prev:
        st.subheader("2. Preview & Edit")
        if not nomor.strip():
            st.info("Isi Nomor Surat Tugas untuk menampilkan preview.")
            return
        tgl_s = tgl_id(tgl)
        auto1 = (f"Menerangkan bahwa dalam rangka melaksanakan perjalanan dinas untuk melaksanakan tugas "
                 f"kedinasan sesuai surat tugas Nomor : {nomor.strip()} pelaksanaan tanggal {tgl_s}, "
                 f"saya benar-benar tidak menggunakan kendaraan dinas.")
        sig = (nomor, tgl_s)
        if ss.get("sp_sig") != sig:
            ss["sp_sig"] = sig
            ss["sp_p1"], ss["sp_p2"] = auto1, SUPER_P2
        with st.expander("✏️ Edit isi surat"):
            st.text_area("Paragraf 1", key="sp_p1", height=110)
            st.text_area("Paragraf 2", key="sp_p2", height=140)
            st.button("↩️ Kembalikan ke teks otomatis", on_click=cb_reset_super_text, key="btn_reset_sp_text")
        d = dict(nama=peg["nama"], nip=peg["nip"], pangkat_gol=pg, jabatan=peg["jabatan"], tanggal=tgl_s,
                 p1=ss["sp_p1"], p2=ss["sp_p2"])
        pdf = build_super_pdf(d, paper)
        for png in render_preview(pdf):
            st.image(png, use_container_width=True)
        st.subheader("3. Download")
        fmt = st.radio("Pilih format download", ["PDF", "Word (.docx)"], horizontal=True, key="sp_fmt")
        base = f"Super_{safe(peg['nama'])}_{tgl.strftime('%Y-%m-%d')}"
        if fmt == "PDF":
            st.download_button("⬇️ Download PDF", pdf, file_name=base + ".pdf", mime="application/pdf",
                               type="primary", key="sp_dl_pdf")
        else:
            st.download_button("⬇️ Download Word", build_super_docx(d, paper), file_name=base + ".docx",
                               mime=DOCX_MIME, type="primary", key="sp_dl_docx")


# ----------------------------------------------------------------------------
# MAIN
# ----------------------------------------------------------------------------
init_state()
with st.sidebar:
    st.header("Menu")
    menu = st.radio("Pilih menu", ["Laporan Perjalanan Translok", "Generate dari Excel",
                                   "Surat Pernyataan Tidak Menggunakan Kendaraan Dinas (Super)"],
                    label_visibility="collapsed")
    st.divider()
    paper = st.selectbox("Ukuran kertas", list(PAPER), index=0)
    use_kop = st.checkbox("Pakai kop surat di laporan", value=True)
    st.subheader("Ukuran & warna foto dokumentasi")
    st.radio("Warna foto", ["Berwarna", "Grayscale (hitam-putih)"], horizontal=True, key="photo_color")
    photo_w = st.slider("Lebar foto (cm)", 3.0, 16.0, 8.5, 0.5)
    photo_h = st.slider("Tinggi maksimal foto (cm)", 3.0, 20.0, 9.5, 0.5,
                        help="Foto yang terlalu tinggi otomatis diperkecil proporsional.")
    if (use_kop and not KOP_PATH.exists()) or not QR_PATH.exists():
        st.warning("Folder 'assets' (kop.png & qr.png) belum ada di repo.")

if menu.startswith("Laporan"):
    page_laporan(paper, use_kop, photo_w, photo_h)
elif menu.startswith("Generate"):
    page_excel(paper, use_kop, photo_w, photo_h)
else:
    page_super(paper)
